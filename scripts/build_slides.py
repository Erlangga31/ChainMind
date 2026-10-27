"""Build reproducible Indonesian and English ChainMind pitch decks."""

from __future__ import annotations

# Slide copy is authored as complete sentences to preserve readable editorial text.
# ruff: noqa: E501
import argparse
from dataclasses import dataclass
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(15, 35, 55)
TEAL = RGBColor(19, 133, 124)
PALE = RGBColor(235, 245, 244)
WHITE = RGBColor(255, 255, 255)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


@dataclass(frozen=True)
class SlideContent:
    title: str
    bullets: tuple[str, ...]
    notes: str
    image: str | None = None


def content(language: str) -> tuple[SlideContent, ...]:
    if language == "id":
        return (
            SlideContent(
                "ChainMind",
                (
                    "Autonomous Disruption Resolver",
                    "Membantu tim manufaktur merespons gangguan pasokan dengan cepat dan terukur.",
                ),
                "Pembuka: ChainMind mengubah sinyal gangguan pasokan menjadi rekomendasi yang dapat ditinjau manusia. Demo berjalan dengan fixture dan sandbox lokal.",
            ),
            SlideContent(
                "Masalah: buffer habis sebelum respons tiba",
                (
                    "Kapal Microchip X terlambat 5 hari.",
                    "Stok hanya mencukupi 3 hari produksi.",
                    "Pencarian vendor, analisis stok, dan keputusan sering tersebar lintas sistem.",
                ),
                "Jelaskan risiko dua hari stockout sebagai hasil fixture demo, bukan prediksi produksi.",
            ),
            SlideContent(
                "Mengapa AWS dan SAP",
                (
                    "SAP S/4HANA dan Ariba berperan sebagai sumber data operasional.",
                    "AWS menyediakan event, compute, orkestrasi, penyimpanan audit, identitas, dan opsi agentic AI.",
                    "Integrasi produksi melalui BTP dan Integration Suite adalah target, bukan koneksi yang aktif di demo.",
                ),
                "Tekankan pembagian system-of-record (SAP) dan system-of-action (ChainMind).",
            ),
            SlideContent(
                "Arsitektur dan alur",
                (
                    "Event → Monitoring → Logistics → Procurement → Orchestrator.",
                    "Step Functions menunggu keputusan manusia sebelum eksekusi.",
                    "Python menghitung coverage, biaya dan ranking; LLM hanya merangkum alasan.",
                ),
                "Bacakan alur kiri ke kanan pada diagram; semua adapter yang berjalan di demo bersifat mock/sandbox.",
                "docs/screenshots/flow-id.png",
            ),
            SlideContent(
                "Empat agen, tanggung jawab jelas",
                (
                    "Monitoring: klasifikasi severity dan persist event.",
                    "Logistics: coverage, stockout, dampak BOM dan estimasi loss.",
                    "Procurement: bandingkan minimal dua vendor.",
                    "Orchestrator: evaluasi ekonomi, rekomendasi dan draf PO.",
                ),
                "Rumus finansial bersifat deterministik; model bahasa tidak menghitung angka.",
            ),
            SlideContent(
                "Demo langsung: dari cuaca ke draf PO",
                (
                    "Microchip X: terlambat 5 hari; buffer 3 hari.",
                    "Vendor B dipilih: lead time 2 hari.",
                    "5.000 pcs × $27 = $135.000; terms Air Freight.",
                    "Status tetap PENDING_APPROVAL.",
                ),
                "Jalankan scenario cuaca, tunjukkan Flow, vendor dan draf PO; jangan mengklaim transaksi SAP nyata.",
                "docs/screenshots/dashboard-id-light.png",
            ),
            SlideContent(
                "Human-in-the-loop dan audit",
                (
                    "Approver meninjau rekomendasi dan mengonfirmasi keputusan.",
                    "Tidak ada PO yang dieksekusi sebelum persetujuan eksplisit.",
                    "Audit trail mencatat delapan langkah; perubahan status dan tracking tetap dapat ditelusuri.",
                ),
                "Tekankan bahwa approve hanya dilakukan di simulator lokal pada demo ini.",
                "docs/screenshots/notification-feed-id.png",
            ),
            SlideContent(
                "Model bisnis — hipotesis ilustratif",
                (
                    "Starter / Professional / Enterprise berdasarkan volume disrupsi dan jumlah agen.",
                    "Biaya onboarding satu kali dapat mencakup integrasi SAP.",
                    "Harga, willingness-to-pay, dan listing Marketplace belum divalidasi.",
                ),
                "Harga pada UI adalah estimasi ilustratif, bukan penawaran atau invoice.",
            ),
            SlideContent(
                "KPI — target ilustratif",
                (
                    "Waktu respons: target menit, dibanding baseline manual yang diasumsikan.",
                    "Line-stoppage risk, beban procurement, dan konsistensi keputusan perlu divalidasi pilot.",
                    "Hasil dashboard adalah fixture/simulasi, bukan dampak pelanggan terukur.",
                ),
                "Jangan menyampaikan nilai KPI sebagai hasil produksi atau forecast.",
            ),
            SlideContent(
                "Roadmap dan risiko",
                (
                    "Demo lokal → pilot sandbox berizin → integrasi tenant → deployment produksi.",
                    "Risiko utama: freshness data, kualitas vendor, akses SAP/AWS, keamanan dan perubahan API.",
                    "Mitigasi: validasi kontrak, guardrails, audit, approval wajib dan rollback.",
                ),
                "Produksi hanya setelah tinjauan security, biaya, governance dan sign-off pelanggan.",
            ),
            SlideContent(
                "Verifikasi dan batas klaim",
                (
                    "SAP sandbox simulating S/4HANA OData and Ariba; AWS verified with LocalStack and cdk synth.",
                    "Backend 163 passed / 9 skipped; Vitest 18 passed; Playwright 11 passed.",
                    "Tidak ada deployment AWS atau koneksi tenant SAP nyata.",
                ),
                "Gunakan klaim persis seperti tertulis; LocalStack dan CDK synth bukan verifikasi akun AWS produksi.",
            ),
        )
    return (
        SlideContent(
            "ChainMind",
            (
                "Autonomous Disruption Resolver",
                "Human-approved supply-chain response for manufacturing.",
            ),
            "Opening: ChainMind turns disruption signals into recommendations that a human can review. This demo uses local fixtures and sandboxes.",
        ),
        SlideContent(
            "Problem: the buffer runs out before the response arrives",
            (
                "Microchip X cargo is delayed five days.",
                "Available stock covers only three days of production.",
                "Supplier search, inventory analysis, and decisions are fragmented across systems.",
            ),
            "Explain the two stockout days as a demo-fixture result, not a production forecast.",
        ),
        SlideContent(
            "Why AWS and SAP",
            (
                "SAP S/4HANA and Ariba are the operational data systems of record.",
                "AWS maps event ingestion, compute, orchestration, audit, identity, and agentic AI options.",
                "Production integration through BTP and Integration Suite is a target, not active connectivity in this demo.",
            ),
            "Describe SAP as system of record and ChainMind as a system of action.",
        ),
        SlideContent(
            "Architecture and flow",
            (
                "Event → Monitoring → Logistics → Procurement → Orchestrator.",
                "Step Functions waits for a human decision before execution.",
                "Python computes coverage, cost, and ranking; the LLM only summarizes rationale.",
            ),
            "Walk through the diagram; all adapters exercised in this demo are mock or sandbox implementations.",
            "docs/screenshots/flow-en.png",
        ),
        SlideContent(
            "Four agents, clear responsibilities",
            (
                "Monitoring: classify severity and persist events.",
                "Logistics: calculate coverage, stockout, BOM impact, and estimated loss.",
                "Procurement: compare at least two alternative vendors.",
                "Orchestrator: evaluate economics, recommend, and draft a PO.",
            ),
            "All financial calculations are deterministic; the language model does not compute money.",
        ),
        SlideContent(
            "Live demo: from weather alert to draft PO",
            (
                "Microchip X: five-day delay; three-day buffer.",
                "Vendor B is selected: two-day lead time.",
                "5,000 pcs × $27 = $135,000; Air Freight.",
                "Status remains PENDING_APPROVAL.",
            ),
            "Trigger the weather scenario and show Flow, vendor comparison, and draft PO. Do not claim a real SAP transaction.",
            "docs/screenshots/dashboard-en-light.png",
        ),
        SlideContent(
            "Human-in-the-loop and audit",
            (
                "An approver reviews the recommendation and explicitly confirms a decision.",
                "No PO executes before approval.",
                "An eight-step audit trail and tracking outputs preserve traceability.",
            ),
            "Clarify that approval and execution in this demo occur only in a local simulator.",
            "docs/screenshots/notification-feed.png",
        ),
        SlideContent(
            "Business model — illustrative hypothesis",
            (
                "Starter / Professional / Enterprise tiers by disruption volume and active agents.",
                "One-time onboarding may cover SAP integration.",
                "Pricing, willingness to pay, and Marketplace listings are not validated.",
            ),
            "UI prices are illustrative estimates, not a quote or invoice.",
        ),
        SlideContent(
            "KPIs — illustrative targets",
            (
                "Response time: a minutes-level target versus an assumed manual baseline.",
                "Line-stoppage risk, procurement workload, and decision consistency require pilot validation.",
                "Dashboard results are fixtures/simulation, not measured customer outcomes.",
            ),
            "Never present KPI values as production results or forecasts.",
        ),
        SlideContent(
            "Roadmap and risks",
            (
                "Local demo → authorized sandbox pilot → tenant integration → production deployment.",
                "Key risks: data freshness, supplier quality, AWS/SAP access, security, and API changes.",
                "Mitigations: contract validation, guardrails, audit, required approval, and rollback.",
            ),
            "Production requires security, cost, governance, and customer sign-off.",
        ),
        SlideContent(
            "Verification and honest scope",
            (
                "SAP sandbox simulating S/4HANA OData and Ariba; AWS verified with LocalStack and cdk synth.",
                "Backend 163 passed / 9 skipped; Vitest 18 passed; Playwright 11 passed.",
                "No AWS deployment or live SAP tenant connection was performed.",
            ),
            "Use the exact claim shown. LocalStack and CDK synthesis do not mean production AWS verification.",
        ),
    )


def add_text(
    slide,
    title: str,
    bullets: tuple[str, ...],
    *,
    has_image: bool,
) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.9), Inches(1.0))
    title_frame = title_box.text_frame
    title_frame.clear()
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.text = title
    title_paragraph.font.name = "Aptos Display"
    title_paragraph.font.size = Pt(30)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = NAVY

    body_width = Inches(6.3 if has_image else 11.4)
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.7), body_width, Inches(4.9))
    frame = body.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {bullet}"
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(23)
        paragraph.font.color.rgb = NAVY
        paragraph.space_after = Pt(18)


def build_deck(language: str, output: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    blank = presentation.slide_layouts[6]
    for index, item in enumerate(content(language), start=1):
        slide = presentation.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = PALE if index % 2 == 0 else WHITE
        band = slide.shapes.add_shape(1, 0, 0, SLIDE_WIDTH, Inches(0.18))
        band.fill.solid()
        band.fill.fore_color.rgb = TEAL
        band.line.fill.background()
        add_text(slide, item.title, item.bullets, has_image=item.image is not None)
        if item.image:
            image_path = ROOT / item.image
            if image_path.is_file():
                with Image.open(image_path) as image:
                    aspect = image.width / image.height
                image_width = min(5.0, 4.5 * aspect)
                image_height = image_width / aspect
                slide.shapes.add_picture(
                    str(image_path),
                    Inches(7.65 + (5.0 - image_width) / 2),
                    Inches(2.15 + (4.5 - image_height) / 2),
                    width=Inches(image_width),
                    height=Inches(image_height),
                )
        footer = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(11.9), Inches(0.25))
        footer.text_frame.text = f"CHAINMIND  ·  {index:02d}"
        footer.text_frame.paragraphs[0].font.size = Pt(10)
        footer.text_frame.paragraphs[0].font.color.rgb = TEAL
        slide.notes_slide.notes_text_frame.text = item.notes
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "docs" / "slides")
    arguments = parser.parse_args()
    for language, filename in (
        ("id", "ChainMind_Pitch_ID.pptx"),
        ("en", "ChainMind_Pitch_EN.pptx"),
    ):
        path = build_deck(language, arguments.output_dir / filename)
        print(f"Built {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
