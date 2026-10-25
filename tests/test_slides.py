from pathlib import Path

from pptx import Presentation

from scripts.build_slides import build_deck


def test_builds_bilingual_pitch_decks_with_eleven_slides_and_speaker_notes(
    tmp_path: Path,
) -> None:
    for language, filename in (
        ("id", "ChainMind_Pitch_ID.pptx"),
        ("en", "ChainMind_Pitch_EN.pptx"),
    ):
        output = build_deck(language, tmp_path / filename)
        deck = Presentation(output)

        assert len(deck.slides) == 11
        assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in deck.slides)
        assert (tmp_path / filename).stat().st_size > 20_000
